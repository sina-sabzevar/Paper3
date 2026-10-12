# -*- coding: utf-8 -*-
"""Thin helpers for building a Persian, right-to-left .pptx with python-pptx.

Everything is measured in POINTS on a 960 x 540 pt canvas (13.333 x 7.5 in),
which is the 1920x1080 design divided by two.

Two things python-pptx will not do on its own and PowerPoint needs:
  * `rtl="1"` on every paragraph, or Persian punctuation and any Latin or
    mixed run inside the line comes out in the wrong order;
  * a `<a:cs>` typeface on every run, because Persian is complex script and
    PowerPoint picks the theme's CS font, not the latin one, for it.
"""
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
import os
from PIL import Image

ICONS = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     "assets", "icons")


def icon(slide, name, cx, cy, h):
    """one of the house icons, centred on (cx, cy) at height h in points."""
    path = os.path.join(ICONS, name + ".png")
    iw, ih = Image.open(path).size
    w = h * iw / float(ih)
    return slide.shapes.add_picture(path, P(cx - w / 2.0), P(cy - h / 2.0),
                                    P(w), P(h))


def _no_theme_style(shape):
    sp = shape._element
    st = sp.find(qn("p:style"))
    if st is not None:
        sp.remove(st)


FONT = "Tahoma"                 # ships with Windows, covers Persian
W, H  = 960.0, 540.0            # the canvas, in points

INK   = RGBColor(0x0B, 0x0B, 0x0B)
INK2  = RGBColor(0x4E, 0x4D, 0x4A)
INK3  = RGBColor(0x86, 0x85, 0x7F)
PLANE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF6, 0xF6, 0xF3)
LINE  = RGBColor(0xE2, 0xE1, 0xDB)
RULE  = RGBColor(0xE8, 0xE8, 0xE4)
RULE2 = RGBColor(0xD4, 0xD3, 0xCC)   # a table rule has to survive projection
S1    = RGBColor(0x2A, 0x78, 0xD6)
CRIT  = RGBColor(0xD0, 0x3B, 0x3B)
MDLBG = RGBColor(0xDB, 0xE6, 0xF2)   # s1 13 pct over panel
OUTBG = RGBColor(0xE4, 0xEB, 0xF3)   # s1  9 pct over panel
OUTLN = RGBColor(0x8F, 0xB2, 0xD9)   # s1 45 pct over line
BAND3 = RGBColor(0xFB, 0xED, 0xED)   # crit 9 pct over white

def P(v):
    return Emu(int(round(v * 12700)))


EMOJI = "Segoe UI Emoji"


def _run(para, text, size, bold=False, color=INK, face=None):
    r = para.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    face = face or FONT
    r.font.name = face
    r.font.color.rgb = color
    rPr = r._r.get_or_add_rPr()
    rPr.set("lang", "fa-IR")
    for tag in ("a:cs", "a:ea"):                 # complex script and east asian
        el = rPr.makeelement(qn(tag), {"typeface": face})
        rPr.append(el)
    return r


def _para(tf, i, align):
    """paragraph i of a text frame, created if needed, set to RTL."""
    while len(tf.paragraphs) <= i:
        tf.add_paragraph()
    p = tf.paragraphs[i]
    pPr = p._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    pPr.set("algn", align)
    return p


def text(slide, x, y, w, h, lines, align="ctr", anchor=MSO_ANCHOR.MIDDLE,
         wrap=False, line_spacing=None):
    """lines: list of (text, size, bold, color) - one paragraph each."""
    tb = slide.shapes.add_textbox(P(x), P(y), P(w), P(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, spec in enumerate(lines):
        p = _para(tf, i, align)
        if line_spacing:
            p.line_spacing = line_spacing
        for sub in (spec if isinstance(spec[0], (list, tuple)) else [spec]):
            _run(p, *sub)
    return tb


def rect(slide, x, y, w, h, fill, stroke=None, stroke_w=0.75, radius=0.07):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        P(x), P(y), P(w), P(h))
    if radius:
        sh.adjustments[0] = radius
    sh.shadow.inherit = False                    # the theme shadow is not wanted
    _no_theme_style(sh)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if stroke is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = stroke
        sh.line.width = Pt(stroke_w)
    sh.text_frame.word_wrap = False
    return sh


def oval(slide, x, y, w, h, fill, stroke=None, stroke_w=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, P(x), P(y), P(w), P(h))
    sh.shadow.inherit = False
    _no_theme_style(sh)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if stroke is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = stroke; sh.line.width = Pt(stroke_w)
    sh.text_frame.word_wrap = False
    return sh


def line(slide, x1, y1, x2, y2, color=LINE, w=0.75, dash=None, arrow=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    P(x1), P(y1), P(x2), P(y2))
    cn.line.color.rgb = color
    cn.line.width = Pt(w)
    cn.shadow.inherit = False
    _no_theme_style(cn)
    ln = cn.line._get_or_add_ln()
    if dash:
        d = ln.makeelement(qn("a:prstDash"), {"val": dash})
        ln.append(d)
    if arrow:
        t = ln.makeelement(qn("a:tailEnd"), {"type": "triangle",
                                             "w": "med", "len": "med"})
        ln.append(t)
    return cn


# ---------------------------------------------------------------- tables
def strip_table_style(tbl):
    """python-pptx stamps every new table with a banded blue style."""
    tblPr = tbl.find(qn("a:tblPr"))
    if tblPr is None:
        return
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        if attr in tblPr.attrib:
            del tblPr.attrib[attr]
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        tblPr.remove(sid)


_EDGES = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]

def cell_borders(cell, edges, color, w=0.75):
    """edges: any of 'L','R','T','B' get a line in `color`; every other edge
    is switched OFF, because an absent border element means inherit and the
    renderer's default for a table is a full grid."""
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in _EDGES:                            # start from a clean slate
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    for tag in reversed(_EDGES):
        on = tag[-1] in edges
        ln = tcPr.makeelement(qn(tag), {"w": str(int(w * 12700)), "cap": "flat",
                                        "cmpd": "sng", "algn": "ctr"})
        if on:
            fill = ln.makeelement(qn("a:solidFill"), {})
            fill.append(ln.makeelement(qn("a:srgbClr"), {"val": str(color)}))
        else:
            fill = ln.makeelement(qn("a:noFill"), {})
        ln.append(fill)
        tcPr.insert(0, ln)                        # lnL, lnR, lnT, lnB come first


def cell_fill(cell, color):
    if color is None:
        cell.fill.background()
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = color


def cell_text(cell, lines, align="ctr", anchor=MSO_ANCHOR.MIDDLE,
              pad=(2, 3, 2, 3), wrap=False):
    tf = cell.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    cell.margin_top, cell.margin_right = P(pad[0]), P(pad[1])
    cell.margin_bottom, cell.margin_left = P(pad[2]), P(pad[3])
    cell.vertical_anchor = anchor
    for i, spec in enumerate(lines):
        p = _para(tf, i, align)
        for sub in (spec if isinstance(spec[0], (list, tuple)) else [spec]):
            _run(p, *sub)
